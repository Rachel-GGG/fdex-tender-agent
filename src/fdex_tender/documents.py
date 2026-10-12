"""Editable personnel chapter and an independent review of actual OOXML."""
import re
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches
from .core import InputError, digest_file, digest_bytes, envelope, issue, require_same_context, safe_path


def assemble(ctx, tender, matrix, index, matched, run_dir):
    require_same_context(ctx, tender, matrix, index, matched)
    problems = list(matrix["issues"]) + list(matched["issues"])
    result = {"artifact_status": "NOT_GENERATED", "files": [], "trace": [], "selected_person_ids": [], "selection_digest": matched["data"]["selection_digest"]}
    if not matched["data"]["selected"] or matched["data"]["confirmation"]["status"] != "CONFIRMED" or any(i["severity"] == "BLOCKER" for i in problems):
        return envelope(ctx, "M13", result, matrix["input_refs"], problems,
                        status="BLOCKED" if any(i["severity"] == "BLOCKER" for i in problems) else "REVIEW")
    manifest = matrix["data"]["template"]
    template = safe_path(ctx["input_root"], manifest["relative_path"])
    if digest_file(template) != manifest["sha256"]:
        problems.append(issue("TEMPLATE_HASH_MISMATCH", manifest["template_id"], "装配时模板哈希已变化"))
        return envelope(ctx, "M13", result, matrix["input_refs"], problems)
    doc = Document(template)
    replacements = {"{{PROJECT_NAME}}": ctx["project_name"], "{{PROJECT_ID}}": ctx["project_id"], "{{LOT_ID}}": ctx["lot_id"]}
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for key, value in replacements.items():
                run.text = run.text.replace(key, value)
    if len(doc.tables) != 1 or [cell.text for cell in doc.tables[0].rows[0].cells] != manifest["columns"]:
        return envelope(ctx, "M13", result, matrix["input_refs"], [issue("TEMPLATE_TABLE_MISMATCH", "template", "模板表格与字段映射不一致")])
    table = doc.tables[0]
    assets = {a["asset_id"]: a for a in index["data"]["assets"]}
    trace = []
    for ordinal, person in enumerate(matched["data"]["selected"], 1):
        asset = assets[person["asset_id"]]
        image_path = safe_path(ctx["input_root"], asset["relative_path"])
        if not image_path.is_file() or digest_file(image_path) != asset["sha256"]:
            return envelope(ctx, "M13", result, matrix["input_refs"], [issue("ASSET_HASH_CHANGED", asset["asset_id"], "确认后证书原件发生变化，需重新确认")])
        cells = table.add_row().cells
        values = [person["role"], person["name"], asset["certificate_id"], asset["valid_until"], asset["relative_path"]]
        for cell, value in zip(cells, values):
            cell.text = value
        doc.add_page_break()
        doc.add_heading(f"人员证书 {ordinal}", level=1)
        doc.add_paragraph(f"{person['name']}  {asset['certificate_id']}")
        doc.add_paragraph("教学示例证书，仅用于合成材料验证。")
        doc.add_picture(str(image_path), width=Inches(6.1))
        trace.append({"section_id": "SECTION-PERSONNEL", "field_id": f"PERSONNEL-ROW-{ordinal}",
                      "person_id": person["person_id"], "asset_id": asset["asset_id"],
                      "attachment_id": f"ATT-{person['asset_id']}", "image_sha256": asset["sha256"],
                      "requirement_ids": [r["requirement_id"] for r in tender["data"]["requirements"] if r["in_personnel_scope"]],
                      "output_file": "personnel-chapter.docx", "output_location": f"人员表第{ordinal}行及证书附件{ordinal}"})
    output = Path(run_dir) / "personnel-chapter.docx"
    doc.core_properties.author = "FDEX Teaching Baseline"
    doc.core_properties.title = "人员资质章节"
    doc.core_properties.subject = "DEMO personnel chapter only"
    doc.save(output)
    result.update(artifact_status="PARTIAL_DRAFT", files=[{"relative_path": output.name, "sha256": digest_file(output), "editable": True}],
                  trace=trace, selected_person_ids=[p["person_id"] for p in matched["data"]["selected"]],
                  template_id=manifest["template_id"], template_version=manifest["version"])
    return envelope(ctx, "M13", result, matrix["input_refs"] + index["input_refs"], problems)


def review_document(ctx, matrix, index, matched, assembled, run_dir):
    require_same_context(ctx, matrix, index, matched, assembled)
    problems, checks = [], []
    coverage = matrix["data"]["coverage"]
    result = {"review_status": "DRAFT", "file_sha256": None, "checks": checks, "coverage": coverage}
    if not assembled["data"]["files"]:
        return envelope(ctx, "M14", result, assembled["input_refs"], [issue("NO_DOCUMENT", "document", "没有可审查的实际Word，前序确认或材料问题待处理")])
    path = safe_path(run_dir, assembled["data"]["files"][0]["relative_path"])
    if not path.exists():
        return envelope(ctx, "M14", result, assembled["input_refs"], [issue("DOCUMENT_MISSING", "document", "实际Word不存在")])
    result["file_sha256"] = digest_file(path)
    if result["file_sha256"] != assembled["data"]["files"][0]["sha256"]:
        problems.append(issue("DOCUMENT_HASH_CHANGED", "document", "Word与装配记录版本不一致"))
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8")
            media = {digest_bytes(z.read(name)) for name in z.namelist() if name.startswith("word/media/")}
        doc = Document(path)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile):
        return envelope(ctx, "M14", result, assembled["input_refs"], [issue("DOCX_UNREADABLE", "document", "实际Word包损坏或关联图片丢失")])
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
    if ctx["project_name"] not in text or ctx["project_id"] not in text or ctx["lot_id"] not in text:
        problems.append(issue("PROJECT_FIELDS_WRONG", "document.title", "实际文档项目字段不一致"))
    if re.search(r"\{\{[^{}]+\}\}|\[TODO[^\]]*\]", text):
        problems.append(issue("PLACEHOLDER_REMAINS", "document", "实际文档仍含未处理占位符"))
    assets = {a["asset_id"]: a for a in index["data"]["assets"]}
    selected = matched["data"]["selected"]
    table_rows = [[c.text for c in row.cells] for row in doc.tables[0].rows[1:]] if len(doc.tables) == 1 else []
    expected_rows = [[p["role"], p["name"], p["certificate_id"], p["valid_until"], assets[p["asset_id"]]["relative_path"]] for p in selected]
    if table_rows != expected_rows:
        problems.append(issue("PERSON_TABLE_WRONG", "PERSONNEL-TABLE", "实际人员表与确认记录不一致，可能包含旧人员或错证书"))
    expected_media = {assets[p["asset_id"]]["sha256"] for p in selected}
    if media != expected_media:
        problems.append(issue("ATTACHMENT_SET_WRONG", "PERSONNEL-ATTACHMENTS", "实际图片缺失、错配或包含旧版本图片"))
    unselected = [p for p in index["data"]["persons"] if p["person_id"] not in {x["person_id"] for x in selected}]
    for person in unselected:
        old_assets = [a for a in index["data"]["assets"] if a["person_id"] == person["person_id"]]
        if person["name"] in text or any(a["certificate_id"] in text or a["sha256"] in media for a in old_assets):
            problems.append(issue("UNSELECTED_PERSON_REMAINS", person["person_id"], "未选人员的姓名、编号或图片残留"))
    rows = matrix["data"]["responses"]
    wanted = {rid for row in rows for rid in row["requirement_ids"]}
    actual = {rid for row in assembled["data"]["trace"] for rid in row["requirement_ids"]}
    if wanted != actual:
        problems.append(issue("RESPONSE_TRACE_MISSING", "trace", "声明范围内要求到实际输出的映射缺失"))
    checks.extend(["actual_docx_opened", "project_fields", "selected_table", "embedded_image_hashes", "no_unselected_person", "no_placeholder", "scope_coverage"])
    result["review_status"] = "CHANGES_REQUIRED" if problems else "DRAFT"
    result["visual_review"] = "NOT_PERFORMED_BY_PROGRAM"
    return envelope(ctx, "M14", result, assembled["input_refs"], problems)
