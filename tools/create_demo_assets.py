#!/usr/bin/env python3
"""Author synthetic teaching fixtures and the declared personnel template."""
import argparse
import shutil
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from lxml import etree
from fdex_tender.core import digest_file, write_json, load_context
from fdex_tender.modules import register, read_tender, material_index, selection_fingerprint

ROOT = Path(__file__).resolve().parents[1]


def make_template(path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    for style_name in ("Normal", "Title", "Heading 1"):
        style = doc.styles[style_name]
        style.font.name = "Noto Sans CJK SC"
        style.font.size = Pt(11 if style_name == "Normal" else 20 if style_name == "Title" else 15)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
    for style in doc.styles:
        props = style.element.find(qn("w:pPr"))
        if props is not None:
            for border in list(props.findall(qn("w:pBdr"))):
                props.remove(border)
    doc.add_paragraph("人员资质章节", style="Title")
    doc.add_paragraph("本章展示经输入确认记录选定的人员与证书。仅覆盖人员资质范围，技术方案、商务报价和全本递交另行准备。")
    doc.add_paragraph("项目名称 {{PROJECT_NAME}}")
    doc.add_paragraph("项目编号 {{PROJECT_ID}}   标包 {{LOT_ID}}")
    doc.add_paragraph("材料为教学示例，不代表任何真实人员资格或企业投标承诺。")
    table = doc.add_table(rows=1, cols=5)
    table.autofit = False
    widths = [0.95, 1.1, 1.55, 1.1, 2.05]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, text, width in zip(table.rows[0].cells, ["角色", "人员", "证书编号", "有效期", "原件位置"], widths):
        cell.text = text
        cell.width = Inches(width)
        props = cell._tc.get_or_add_tcPr()
        shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E9EEF3"); props.append(shade)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            node.set(qn("w:" + key), val)
        borders.append(node)
    table._tbl.tblPr.append(borders)
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    doc.core_properties.author = "FDEX Teaching Baseline"
    # Embed the OFL font, rather than depending on a host font installation.
    # The all-zero font key means obfuscation leaves bytes unchanged (MS-OE376).
    font_table = next(rel.target_part for rel in doc.part.rels.values() if rel.reltype.endswith("/fontTable"))
    tree = etree.fromstring(font_table.blob)
    font = OxmlElement("w:font"); font.set(qn("w:name"), "Noto Sans CJK SC")
    embedded = Part(PackURI("/word/fonts/noto-sc.odttf"), "application/vnd.openxmlformats-officedocument.obfuscatedFont",
                    (ROOT / "assets/fonts/NotoSansCJKsc-Regular.otf").read_bytes(), doc.part.package)
    rid = font_table.relate_to(embedded, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font")
    embed = OxmlElement("w:embedRegular"); embed.set(qn("r:id"), rid)
    embed.set(qn("w:fontKey"), "{00000000-0000-0000-0000-000000000000}")
    embed.set(qn("w:subsetted"), "false"); font.append(embed); tree.append(font)
    font_table._blob = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def make_certificate(path, ordinal):
    image = Image.new("RGB", (1200, 700), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 35)
        title = ImageFont.truetype("DejaVuSans.ttf", 52)
    except OSError:
        font = title = ImageFont.load_default(size=30)
    draw.rectangle((30, 30, 1170, 670), outline="#64748b", width=3)
    draw.text((90, 85), "DEMO CERTIFICATE", font=title, fill="#111827")
    lines = ["TRAINING ONLY - NOT A REAL QUALIFICATION", f"Person ID: PERSON-{ordinal:02}",
             f"Certificate: DEMO-CERT-{ordinal:03}", "Qualification: Demo professional service",
             "Valid until: 2028-12-31", "No seal or official issuer is represented."]
    for i, line in enumerate(lines):
        draw.text((90, 190 + i * 63), line, font=font, fill="#334155")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def build():
    root = ROOT / "samples/e2e/core-normal"
    root.mkdir(parents=True, exist_ok=True)
    write_json(root / "context.json", {"project_id": "DEMO-BID-001", "lot_id": "LOT-001", "project_name": "示例专业服务项目",
                                       "company_id": "DEMO-COMPANY", "source_mode": "DEMO", "analysis_time": "2026-10-13T09:00:00+08:00",
                                       "business_participation": "EXISTING_GO", "note": "所有企业、人员、确认人和证书均虚构；无需联网和API key"})
    tender = """# 示例专业服务采购文件节选
项目编号 DEMO-BID-001  标包 LOT-001
本节选为教学编写，不是公开招标文件；只有下列声明范围。
[R-01][人员] 人员不少于3名，人员表必须列出姓名和本次职责。
[R-02][资格] 每名所选人员必须持有示例专业服务证书。
[R-03][资格] 证书有效期以2026-11-01为基准，必须仍在有效期内。
[R-04][格式] 必须在人员表填写证书编号和有效期。
[R-05][格式] 必须为每名所选人员附清晰的对应证书图片。
[R-06][人员] 必须由资料负责人确认本次选用人员；候选匹配不是正式选用。
[R-07][格式] 必须使用项目提供的人员表版次1.0，生成可编辑Word。
[R-08][格式] 本章必须标明项目名称、项目编号和标包。
[R-09][报价] 全本报价必须说明含税口径；本人员章节不制作报价。
[R-10][技术] 全本必须提供实施计划；本节选只声明人员章节范围。
[R-11][评分] 全本评分包括服务方案；本节选未公布分值，不得补造分数。
[R-12][商务] 全本投标截止为2026-11-01T10:00:00+08:00；没有登记其他前置截止。
"""
    (root / "tender").mkdir(exist_ok=True)
    (root / "tender/tender.md").write_text(tender, encoding="utf-8")
    sources = [{"source_id": "SRC-TENDER", "relative_path": "tender/tender.md", "kind": "tender_markdown", "version": "1.0", "public": True,
                "project_id": "DEMO-BID-001", "lot_id": "LOT-001"}]
    names = ["示例人员甲", "示例人员乙", "示例人员丙", "示例人员丁"]
    for i, name in enumerate(names, 1):
        image_rel = f"materials/cert-{i:02}.png"
        make_certificate(root / image_rel, i)
        sidecar = f"""# 教学人员说明 原件字段由样例作者给定 未进行OCR
person_id: PERSON-{i:02}
name: {name}
asset_id: ASSET-{i:02}
certificate_id: DEMO-CERT-{i:03}
qualification: 示例专业服务
valid_until: 2028-12-31
image: {image_rel}
company_id: DEMO-COMPANY
"""
        (root / f"materials/person-{i:02}.md").write_text(sidecar, encoding="utf-8")
        sources.extend([{"source_id": f"SRC-PERSON-{i:02}", "relative_path": f"materials/person-{i:02}.md", "kind": "person_sidecar", "version": "1.0", "public": True},
                        {"source_id": f"SRC-CERT-{i:02}", "relative_path": image_rel, "kind": "certificate_image", "version": "1.0", "public": True}])
    make_template(ROOT / "templates/personnel-v1.docx")
    (root / "template").mkdir(exist_ok=True)
    shutil.copy2(ROOT / "templates/personnel-v1.docx", root / "template/personnel-v1.docx")
    write_json(root / "template-manifest.json", {"template_id": "TPL-PERSONNEL", "version": "1.0", "relative_path": "template/personnel-v1.docx",
                                                 "sha256": digest_file(root / "template/personnel-v1.docx"), "columns": ["角色", "人员", "证书编号", "有效期", "原件位置"]})
    sources.append({"source_id": "SRC-TEMPLATE", "relative_path": "template/personnel-v1.docx", "kind": "template", "version": "1.0", "public": True})
    write_json(root / "sources.json", {"sources": sources})
    ctx = load_context(root, "FIXTURE")
    registered = register(ctx); tender_result = read_tender(ctx, registered); index = material_index(ctx, registered)
    ids = ["PERSON-01", "PERSON-02", "PERSON-03"]
    roles = {"PERSON-01": "负责人", "PERSON-02": "业务分析", "PERSON-03": "交付测试"}
    write_json(root / "selection.json", {"project_id": ctx["project_id"], "lot_id": ctx["lot_id"], "source_mode": "DEMO",
                                         "status": "CONFIRMED", "person_ids": ids, "roles": roles, "target_digest": selection_fingerprint(ctx, tender_result, index, ids, roles),
                                         "confirmed_by": "DEMO-资料负责人", "confirmed_at": "2026-10-13T09:05:00+08:00", "record_kind": "SYNTHETIC_INPUT_CONFIRMATION"})
    for scenario in ("unconfirmed", "replacement", "expired", "missing-image"):
        target = ROOT / "samples/e2e" / scenario
        if target.exists():
            raise ValueError(f"Fixture directory already exists: {target}")
        shutil.copytree(root, target)
        if scenario == "unconfirmed":
            (target / "selection.json").unlink()
        elif scenario == "replacement":
            ctx = load_context(target, "FIXTURE")
            reg = register(ctx); t = read_tender(ctx, reg); idx = material_index(ctx, reg)
            ids = ["PERSON-01", "PERSON-02", "PERSON-04"]
            roles = {"PERSON-01": "负责人", "PERSON-02": "业务分析", "PERSON-04": "交付测试"}
            record = __import__("json").loads((target / "selection.json").read_text())
            record.update(person_ids=ids, roles=roles, target_digest=selection_fingerprint(ctx, t, idx, ids, roles), confirmed_at="2026-10-13T10:05:00+08:00")
            write_json(target / "selection.json", record)
        elif scenario == "expired":
            path = target / "materials/person-03.md"
            path.write_text(path.read_text().replace("2028-12-31", "2025-12-31"), encoding="utf-8")
        else:
            (target / "materials/cert-03.png").unlink()
    print("Created bounded public DEMO fixtures and editable template")


if __name__ == "__main__":
    build()
